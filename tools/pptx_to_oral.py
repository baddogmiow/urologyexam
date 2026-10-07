#!/usr/bin/env python3
"""把口試簡報 (.pptx) 轉成口試練習頁的匯入格式 (Q:/A:)。

用法：
  python tools/pptx_to_oral.py 簡報1.pptx 簡報2.pptx ... > oral.txt
  python tools/pptx_to_oral.py 簡報資料夾/ > oral.txt

辨識規則（依常見簡報結構）：
  - 首行以「Case」開頭的投影片 = 病例敘述，會附在該 Case 之後每一題前面
  - 首行為「第N題」的投影片 = 題目
  - 首行為「第N題: 標題」的投影片 = 詳解，「Checkpoint:」之後是給分點，
    其餘短行（如 Campbell 12th edition、文獻引用）當成出處
  - 圖片不轉換，只標示張數，請回簡報對照
沒有標出處的題目會標「待查」。輸出後請抽查幾題與原簡報是否一致。
"""
import os
import re
import sys

from pptx import Presentation

Q_RE = re.compile(r'^第\s*([0-9一二三四五六七八九十]+)\s*題\s*$')
A_RE = re.compile(r'^第\s*([0-9一二三四五六七八九十]+)\s*題\s*[:：]')
CASE_RE = re.compile(r'^\s*case\b', re.I)


def slide_text(slide):
    lines, pics = [], 0
    for sh in slide.shapes:
        if sh.shape_type == 13:
            pics += 1
        if sh.has_text_frame:
            lines += [l.strip() for l in sh.text_frame.text.splitlines() if l.strip()]
        if getattr(sh, 'has_table', False) and sh.has_table:
            for r in sh.table.rows:
                lines.append(' | '.join(c.text.strip() for c in r.cells))
    return lines, pics


def convert(path):
    prs = Presentation(path)
    case_label, stem = '', ''
    items, order = {}, []
    for slide in prs.slides:
        lines, pics = slide_text(slide)
        if not lines:
            continue
        head = lines[0]
        if CASE_RE.match(head):
            case_label = head
            stem = ' '.join(lines[1:])
            continue
        m = Q_RE.match(head)
        if m:
            key = (case_label, m.group(1))
            items[key] = {'q': ' '.join(lines[1:]), 'ans': [], 'cp': [], 'src': [], 'pics': 0}
            order.append(key)
            items[key]['stem'] = stem
            continue
        m = A_RE.match(head)
        if m and order:
            key = (case_label, m.group(1))
            it = items.get(key) or items[order[-1]]
            it['pics'] += pics
            body, in_cp = lines[1:], False
            for l in body:
                if re.match(r'^checkpoint\s*[:：]?', l, re.I):
                    in_cp = True
                    rest = re.sub(r'^checkpoint\s*[:：]?\s*', '', l, flags=re.I)
                    if rest:
                        it['cp'].append(rest)
                elif in_cp and len(l) > 60 and not re.search(r'doi|PMID|edition', l, re.I):
                    it['cp'].append(l)
                elif re.search(r'doi|PMID|edition|guideline|et al', l, re.I) and len(l) < 400:
                    it['src'].append(l)
                elif in_cp:
                    it['cp'].append(l)
                else:
                    it['ans'].append(l)
            continue
        # 續頁（只有圖或出處）：併入最近一題
        if order:
            it = items[order[-1]]
            it['pics'] += pics
            for l in lines:
                if re.search(r'doi|PMID|edition|guideline|et al', l, re.I):
                    it['src'].append(l)
    out = []
    for key in order:
        it = items[key]
        label = re.sub(r'\s+', ' ', key[0]).strip()
        name = f"【{os.path.splitext(os.path.basename(path))[0]} {label} 第{key[1]}題】"
        q = (it['stem'] + ' ' if it['stem'] else '') + it['q']
        a = ['【詳解】'] + (it['ans'] or ['（簡報未提供文字詳解）'])
        if it['cp']:
            a += ['【給分點】' + ' '.join(it['cp'])]
        a += ['【出處】' + ('；'.join(dict.fromkeys(it['src'])) if it['src'] else '待查（簡報未標示）')]
        if it['pics']:
            a += [f"【圖片】詳解頁含 {it['pics']} 張圖，請回簡報對照"]
        out.append('Q: ' + name + q + '\nA: ' + '\n'.join(a))
    return out


def main(args):
    paths = []
    for a in args:
        if os.path.isdir(a):
            paths += sorted(os.path.join(a, f) for f in os.listdir(a) if f.lower().endswith('.pptx'))
        else:
            paths.append(a)
    if not paths:
        sys.exit(__doc__)
    total = []
    for p in paths:
        qs = convert(p)
        sys.stderr.write(f'{os.path.basename(p)}: {len(qs)} 題\n')
        total += qs
    sys.stderr.write(f'共 {len(total)} 題\n')
    sys.stdout.reconfigure(encoding='utf-8')
    print('\n\n'.join(total))


if __name__ == '__main__':
    main(sys.argv[1:])
