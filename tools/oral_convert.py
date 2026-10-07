#!/usr/bin/env python3
"""把口試題目（PPT 或 PDF）轉成口試練習頁的匯入格式。PPT 與 PDF 可以放在同一個資料夾。

一個 Case（一大題）一筆，格式與網頁相同：
  Q: 【領域｜Case 名稱】病例：... （第一題）題目 （第二題）題目 ...
  A: ■ 第一題
     詳解...
     【給分點】...
     【出處】...

用法：
  python tools/oral_convert.py 資料夾 --domain "01 結石" -o oral_結石.txt
  python tools/oral_convert.py a.pptx b.pdf --domain "01 結石" -o oral_結石.txt

領域：01 結石、02 前列腺疾病含癌症、03 腎臟輸尿管膀胱腫瘤、04 腎上腺&外生殖器腫瘤、
      06 排尿功能障礙&女性泌尿、07 男性學、08 小兒泌尿、09 移植、10 感染、11 外傷
不同領域請分開執行。請用 -o 寫檔（Windows PowerShell 的 > 會讓中文變亂碼）。

輸入類型（自動判斷）：
  .pptx                    投影片：Case / 第N題 / 第N題: 解答（含 Checkpoint、Ref.）
  .pdf（投影片轉出的 PDF） 每頁一張投影片，規則同上
  .pdf（文件 / 共筆）       「第N題」加「Answer：」或表格，版式不固定，準確度較低
圖片不轉換，只標示張數。轉完會印出「需要核對」的清單，請務必對照原檔抽查。
"""
import argparse
import os
import re
import sys

NUM = '一二三四五六七八九十'
GLYPH = '❖•●○◆·▪■□➢➤⚫'
QM = re.compile(r'^第\s*([0-9一二三四五六七八九十]+)\s*題\s*$')
AM = re.compile(r'^第\s*([0-9一二三四五六七八九十]+)\s*題\s*[:：]')
QINLINE = re.compile(r'^第\s*([0-9一二三四五六七八九十]+)\s*題\s*[:：]?\s*(.*)$')
CASE = re.compile(r'^\s*case\b', re.I)
SRC = re.compile(r'^\s*ref\b|doi|PMID|edition|guideline|et al', re.I)
CP = re.compile(r'^\s*checkpoint\s*[:：]?', re.I)
NOISE = re.compile(r'口試出題範例|^for考官$')
STEM_START = re.compile(r'^(Case\b|\(?（?\d+[)）]\s*115年口試|第[一二三四五六七八九十]波|\d{2,3}\s*歲|一位|一名|一對|一個\d+個月)')


def cn(n):
    n = str(n)
    return NUM[int(n) - 1] if n.isdigit() and 1 <= int(n) <= 10 else n


def num(n):
    n = str(n)
    return int(n) if n.isdigit() else (NUM.index(n) + 1 if n in NUM else 0)


def unglyph(s):
    return s.lstrip(GLYPH + ' \t').strip()


def bullet(l):
    """行首的 ❖、• 等符號統一成「• 」，沒有符號的行不動。"""
    l = l.strip()
    if l and l[0] in GLYPH:
        u = unglyph(l)
        return u if re.match(r'^\d+[.)、．]', u) else '• ' + u
    return l


def new_case(label='', stem=''):
    return {'label': label, 'stem': stem, 'subs': [], 'flags': []}


def new_sub(n, q=''):
    return {'num': cn(n), 'q': q, 'ans': [], 'cp': [], 'src': [], 'pics': 0, 'pages': []}


def add_answer_line(sub, l, state):
    """把一行放進詳解／給分點／出處。state 是 {'in_cp': bool}"""
    l = l.strip()
    if not l:
        return
    u = unglyph(l)
    if CP.match(u):
        state['in_cp'] = True
        rest = CP.sub('', u).strip()
        if rest:
            sub['cp'].append(rest)
    elif SRC.search(u) and len(u) < 400:
        sub['src'].append(re.sub(r'^ref\.?\s*', '', u, flags=re.I))
        state['in_cp'] = False
    elif state['in_cp']:
        sub['cp'].append(u)
    else:
        sub['ans'].append(bullet(l))


# ------------------------------------------------------------------ 投影片（PPT / 投影片式 PDF）
ANS = re.compile(r'^\s*(?:ans(?:wer)?|解答|答)\s*[:：]\s*(.*)$', re.I)
GRADING = re.compile(r'考官評分|評分重點|評分標準')


def parse_slides(slides):
    """slides: [(lines, pics)]，頁碼 = 順序（從 1 開始）。
    支援兩種投影片寫法：
      A. 題目頁「第N題」＋ 解答頁「第N題: 解答」
      B. 「第N題」頁裡用「Ans:」分隔題目與答案（題目與答案可同頁或分頁）
    沒有標題的續頁併入前一題的詳解；Case 之後、第一題之前的續頁併入病例敘述；
    「考官評分」頁整份附在該 Case 最後一題後面。"""
    cases, cur, sub = [], None, None
    for page_no, (lines, pics) in enumerate(slides, 1):
        lines = [l for l in lines if l.strip() and not NOISE.search(l.strip())]
        if not lines:                             # 純圖片頁
            if sub is not None:
                sub['pics'] += pics
                sub['pages'].append(page_no)
            continue
        head = unglyph(lines[0])
        if CASE.match(head):
            cur = new_case(re.sub(r'\s+', ' ', head).strip(), ' '.join(unglyph(x) for x in lines[1:]))
            cases.append(cur)
            sub = None
            continue
        if cur is None:
            cur = new_case('')
            cases.append(cur)
        if GRADING.search(head):
            cur['grading'] = [unglyph(x) for x in lines[1:]] if len(head) < 6 else [re.sub(r'^考官評分\s*', '', unglyph(lines[0]))] + [unglyph(x) for x in lines[1:]]
            if sub is not None:
                sub['pages'].append(page_no)
            continue
        m = QM.match(head) or AM.match(head)
        if m:
            n = cn(m.group(1))
            body = lines[1:]
            existing = next((x for x in cur['subs'] if x['num'] == n), None)
            # 找「Ans:」分隔行
            split = None
            for i, l in enumerate(body):
                if ANS.match(unglyph(l)):
                    split = i
                    break
            if existing is None:
                sub = new_sub(n, '')
                cur['subs'].append(sub)
                qlines, alines = (body[:split], body[split:]) if split is not None else (body, [])
            else:
                sub = existing
                qlines, alines = (body[:split], body[split:]) if split is not None else ([], body)
                if sub['q'].strip():
                    qlines = []                   # 已經有題目，這頁的前半段不再當題目
            st = {'in_cp': False}
            q = []
            for l in qlines:
                ul = unglyph(l)
                if CP.match(ul) or st['in_cp']:
                    add_answer_line(sub, l, st) if CP.match(ul) else sub['cp'].append(ul)
                    st['in_cp'] = True
                else:
                    q.append(ul)
            if q:
                sub['q'] = (sub['q'] + ' ' + ' '.join(q)).strip()
            st = {'in_cp': False}
            for l in alines:
                mm = ANS.match(unglyph(l))
                if mm:
                    if mm.group(1).strip():
                        add_answer_line(sub, mm.group(1), st)
                    continue
                add_answer_line(sub, l, st)
            sub['pics'] += pics
            if alines or existing is not None:
                sub['pages'].append(page_no)
            continue
        # 沒有標題的續頁
        if not cur['subs']:                       # 病例敘述的續頁
            cur['stem'] = (cur['stem'] + ' ' + ' '.join(unglyph(x) for x in lines)).strip()
            continue
        sub['pics'] += pics
        sub['pages'].append(page_no)
        st = {'in_cp': False}
        for l in lines:
            mm = ANS.match(unglyph(l))
            if mm:
                if mm.group(1).strip():
                    add_answer_line(sub, mm.group(1), st)
                continue
            add_answer_line(sub, l, st)
    return cases


def pptx_slides(path):
    from pptx import Presentation
    out = []
    for slide in Presentation(path).slides:
        lines, pics = [], 0
        for sh in slide.shapes:
            if sh.shape_type == 13:
                pics += 1
            if sh.has_text_frame:
                lines += [l.strip() for l in sh.text_frame.text.splitlines() if l.strip()]
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    lines.append(' | '.join(c.text.strip() for c in r.cells))
        out.append((lines, pics))
    return out


# ------------------------------------------------------------------ PDF 共用
def join_wrapped(text):
    """合併 PDF 的硬換行：項目符號、編號、第N題、Case、Answer 開頭才算新的一行。"""
    start = re.compile(r'^\s*(?:[' + re.escape(GLYPH) + r'\-–]|\d+[.)、．]|[A-Za-z][.)]\s|[①-⑩]|\(\d+\)|第\s*[0-9一二三四五六七八九十]+\s*(?:小)?題|Case\b|Answer\s*[:：]|Ref\b|Checkpoint)', re.I)
    text = re.sub('[\ue000-\uf8ff]', '❖', text)      # Wingdings 等私有字元其實是項目符號
    heading = re.compile(r'^\s*(?:第\s*[0-9一二三四五六七八九十]+\s*(?:小)?題\s*[:：]?|[❖•●○◆·▪■□➢➤⚫\s]*Case\s*[0-9A-Za-z\-]*\s*[:：]?|(?:Ans(?:wer)?|解答)\s*[:：]?|考官評分.*)\s*$', re.I)
    out = []
    for raw in text.split('\n'):
        if not raw.strip():
            continue
        if out and not start.match(raw) and not heading.match(out[-1]):
            out[-1] += raw
        else:
            out.append(raw)
    return [re.sub(r'[ \t]+', ' ', l).strip() for l in out if l.strip()]


def pdf_pages(path):
    import pymupdf
    return pymupdf.open(path)


def is_slide_pdf(doc):
    """每頁字少，且多頁以「第N題」或 Case 開頭 → 投影片式 PDF"""
    chars = sorted(len(p.get_text()) for p in doc)
    median = chars[len(chars) // 2] if chars else 0
    starts = 0
    for p in doc:
        first = next((l.strip() for l in p.get_text().split('\n') if l.strip()), '')
        if QINLINE.match(unglyph(first)) or CASE.match(unglyph(first)):
            starts += 1
    return median < 450 and starts >= max(3, len(doc) // 3)


def pdf_slide_cases(doc):
    slides = []
    for p in doc:
        lines = join_wrapped(p.get_text())
        slides.append((lines, len(p.get_images())))
    return parse_slides(slides)


# ------------------------------------------------------------------ 文件式 PDF（共筆）
def pdf_doc_cases(doc):
    """文件式 PDF：以「Case」、「第N題」、「Answer：」切分；題號從頭開始就視為新的 Case。"""
    events = []
    for pi, page in enumerate(doc):
        trects, rows_all = [], []
        try:
            tabs = page.find_tables().tables
        except Exception:
            tabs = []
        for t in tabs:
            parsed = []
            for r in t.extract():
                cells = []
                for j, c in enumerate(r):
                    if c and c.strip() and (not cells or cells[-1][1] != c.strip()):
                        cells.append((j, c.strip()))
                if cells:
                    parsed.append(cells)
            left = min((c[0][0] for c in parsed if len(c) == 2), default=None)
            started, header = False, []
            for cells in parsed:
                if len(cells) >= 2:
                    c0, c1 = cells[0][1], ' '.join(x[1] for x in cells[1:])
                    if QINLINE.match(c0) and QM.match(c0.split('\n')[0].strip()):
                        started = True
                        events.append(('row', c0, c1))
                    elif started:
                        events.append(('cont', c0, c1))
                    else:
                        header += [x[1] for x in cells]
                else:
                    j, c = cells[0]
                    if QM.match(c.split('\n')[0].strip()):
                        started = True
                        events.append(('row', c, ''))
                    elif started:
                        events.append(('qcont' if (left is not None and j <= left) else 'acont', c, ''))
                    else:
                        header.append(c)
            if header:
                events.append(('text', '\n'.join(header)))
            trects.append(t.bbox)
        rest = []
        for b in page.get_text('blocks'):
            x0, y0, x1, y1, txt = b[:5]
            inside = any(x0 >= tb[0] - 2 and y0 >= tb[1] - 2 and x1 <= tb[2] + 2 and y1 <= tb[3] + 2 for tb in trects)
            if not inside and txt.strip():
                rest.append((y0, txt))
        if rest:
            events.append(('text', ''.join(t for _, t in sorted(rest))))

    cases, cur, sub = [], None, None
    mode = 'stem'

    def restart(n):
        """題號從頭開始：新 Case；把上一個 Case 結尾屬於新病例的段落搬過來。"""
        nonlocal cur, sub, mode
        stem = ''
        if cur and cur['subs']:
            last = cur['subs'][-1]
            buf = last['ans'] if last['ans'] else []
            idx = None
            for i, l in enumerate(buf):
                if STEM_START.match(unglyph(l)):
                    idx = i
            if idx is not None:
                stem = ' '.join(unglyph(x) for x in buf[idx:])
                last['ans'] = buf[:idx]
        cur = new_case('', stem)
        cases.append(cur)
        sub = None
        mode = 'stem'

    def start_sub(n, rest_q):
        nonlocal cur, sub, mode
        if cur is None or (sub is not None and cur['subs'] and num(n) <= num(cur['subs'][-1]['num'])):
            restart(n)
        elif cur is None:
            restart(n)
        sub = new_sub(n, rest_q.strip())
        cur['subs'].append(sub)
        mode = 'q'

    for ev in events:
        kind = ev[0]
        if kind == 'row':
            m = QINLINE.match(ev[1].split('\n')[0].strip())
            qrest = ' '.join(join_wrapped(ev[1].split('\n', 1)[1])) if '\n' in ev[1] else m.group(2)
            start_sub(m.group(1), qrest)
            for l in join_wrapped(ev[2]):
                sub['ans'].append(l)
            mode = 'a'
        elif kind == 'cont' and sub:
            sub['q'] = (sub['q'] + ' ' + ' '.join(join_wrapped(ev[1]))).strip()
            sub['ans'] += join_wrapped(ev[2])
        elif kind == 'qcont' and sub:
            sub['q'] = (sub['q'] + ' ' + ' '.join(join_wrapped(ev[1]))).strip()
        elif kind == 'acont' and sub:
            sub['ans'] += join_wrapped(ev[1])
        elif kind == 'text':
            for l in join_wrapped(ev[1]):
                s = unglyph(l)
                if re.match(r'^第[一二三四五六七八九十]波', s) or re.match(r'^\(?（?\(?\d+[)）]\s*115年口試|^115年口試', s):
                    continue
                if CASE.match(s) and not QINLINE.match(s):
                    cur = new_case(re.sub(r'\s+', ' ', s))
                    cases.append(cur)
                    sub, mode = None, 'stem'
                    continue
                m = QINLINE.match(s)
                if m and re.match(r'^第\s*[0-9一二三四五六七八九十]+\s*題\s*[:：]?', s):
                    start_sub(m.group(1), m.group(2))
                    continue
                if re.match(r'^Answer\s*[:：]', s, re.I):
                    mode = 'a'
                    rest = re.sub(r'^Answer\s*[:：]\s*', '', s, flags=re.I)
                    if sub is not None and rest:
                        sub['ans'].append(rest)
                    continue
                if cur is None:
                    cur = new_case('')
                    cases.append(cur)
                if sub is None:
                    cur['stem'] = (cur['stem'] + ' ' + s).strip()
                elif mode == 'q' and not (l[:1] in GLYPH):
                    sub['q'] = (sub['q'] + ' ' + s).strip()
                else:
                    mode = 'a'
                    sub['ans'].append(bullet(l))
    for c in cases:
        for s in c['subs']:
            s['ans'] = [l for l in s['ans'] if l.strip()]
    return cases


# ------------------------------------------------------------------ 輸出與檢查
def check(case, src):
    f = []
    if not case['subs']:
        f.append('沒有偵測到任何小題')
    if not case['stem'].strip():
        f.append('沒有病例敘述')
    if len(case['subs']) == 1:
        f.append('只有 1 個小題，請確認是否漏題')
    nums = [num(s['num']) for s in case['subs']]
    if nums and nums != list(range(1, len(nums) + 1)):
        f.append('小題編號不連續：' + '、'.join(s['num'] for s in case['subs']))
    for s in case['subs']:
        if not s['q'].strip():
            f.append('第%s題抓不到題目文字' % s['num'])
        if not ' '.join(s['ans']).strip():
            pg = '、'.join(str(p) for p in dict.fromkeys(s['pages']))
            f.append('NOANS|第%s題|%s|%s' % (s['num'], pg, '答案頁含 %d 張圖' % s['pics'] if s['pics'] else ''))
    return f


def render(case, domain, idx, src):
    base = os.path.splitext(os.path.basename(src))[0]
    label = (base + ' ' + case['label']) if case['label'] else ('%s Case %d' % (base, idx))
    head = '【' + (domain + '｜' if domain else '') + label + '】'
    q = head + ('病例：' + case['stem'] + ' ' if case['stem'].strip() else '')
    q += ' '.join('（第%s題）%s' % (s['num'], s['q']) for s in case['subs'])
    blocks = []
    for s in case['subs']:
        a = ['■ 第%s題' % s['num']]
        a += s['ans'] or ['（沒有文字詳解，待補）']
        if s['cp']:
            a.append('【給分點】' + ' '.join(s['cp']))
        a.append('【出處】' + ('；'.join(dict.fromkeys(s['src'])) if s['src'] else '待查（原檔未標示）'))
        if s['pics']:
            a.append('【圖片】詳解頁含 %d 張圖，請回原檔對照' % s['pics'])
        if s is case['subs'][-1] and case.get('grading'):
            a.append('【考官評分】' + ' '.join(case['grading']))
        blocks.append('\n'.join(a))
    return 'Q: ' + q + '\nA: ' + '\n\n'.join(blocks)


def convert_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.pptx':
        return parse_slides(pptx_slides(path)), 'PPT 投影片'
    doc = pdf_pages(path)
    if is_slide_pdf(doc):
        return pdf_slide_cases(doc), 'PDF（投影片式）'
    return pdf_doc_cases(doc), 'PDF（文件式，準確度較低）'


def main():
    ap = argparse.ArgumentParser(description='口試題目（PPT / PDF）-> 口試練習頁匯入格式')
    ap.add_argument('paths', nargs='+', help='.pptx / .pdf 檔或資料夾')
    ap.add_argument('--domain', default='', help='領域，例如 "01 結石"')
    ap.add_argument('-o', '--output', default='', help='輸出檔（UTF-8）')
    args = ap.parse_args()
    paths = []
    for a in args.paths:
        if os.path.isdir(a):
            paths += sorted(os.path.join(a, f) for f in os.listdir(a) if f.lower().endswith(('.pptx', '.pdf')))
        else:
            paths.append(a)
    if not paths:
        sys.exit('找不到 .pptx 或 .pdf 檔')
    out, report, noans, total = [], [], [], 0
    for p in paths:
        try:
            cases, kind = convert_file(p)
        except ImportError as e:
            sys.exit('缺少套件：%s。請執行 pip install python-pptx pymupdf' % e.name)
        cases = [c for c in cases if c['subs']]
        sys.stderr.write('%s [%s]：%d 個 Case、%d 個小題\n' % (os.path.basename(p), kind, len(cases), sum(len(c['subs']) for c in cases)))
        for i, c in enumerate(cases, 1):
            total += 1
            out.append(render(c, args.domain, i, p))
            for fl in check(c, p):
                name = '%s 第 %d 個 Case（%s）' % (os.path.basename(p), i, (c['label'] or c['stem'])[:24])
                if fl.startswith('NOANS|'):
                    _, qn, pg, note = fl.split('|')
                    noans.append('%s %s：原檔第 %s 頁%s' % (name, qn, pg or '？', '（%s）' % note if note else ''))
                else:
                    report.append('%s：%s' % (name, fl))
    sys.stderr.write('共 %d 個 Case\n' % total)
    if not args.domain:
        sys.stderr.write('提醒：沒有指定 --domain，匯入後會被歸為「未分類」。\n')
    if noans:
        sys.stderr.write('\n【完全沒有文字詳解，請手動補上】（共 %d 題）：\n' % len(noans) + '\n'.join('  - ' + r for r in noans) + '\n')
        sys.stderr.write('  輸出檔裡這些小題的詳解是「（沒有文字詳解，待補）」，用記事本搜尋「待補」就能找到。\n')
    if report:
        shown = report[:25]
        sys.stderr.write('\n需要核對（共 %d 項，先列前 %d 項）：\n' % (len(report), len(shown)) + '\n'.join('  - ' + r for r in shown) + '\n')
        if args.output:
            rp = os.path.splitext(args.output)[0] + '_核對清單.txt'
            with open(rp, 'w', encoding='utf-8', newline='\n') as f:
                f.write(('【沒有文字詳解，請手動補上】\n' + '\n'.join(noans) + '\n\n' if noans else '') + '【其他需要核對】\n' + '\n'.join(report) + '\n')
            sys.stderr.write('完整清單已寫入 %s\n' % rp)
    elif not noans:
        sys.stderr.write('自動檢查沒有發現問題（仍請抽查）。\n')
    if noans and not report and args.output:
        rp = os.path.splitext(args.output)[0] + '_核對清單.txt'
        with open(rp, 'w', encoding='utf-8', newline='\n') as f:
            f.write('【沒有文字詳解，請手動補上】\n' + '\n'.join(noans) + '\n')
        sys.stderr.write('清單已寫入 %s\n' % rp)
    text = '\n\n'.join(out) + '\n'
    if args.output:
        with open(args.output, 'w', encoding='utf-8', newline='\n') as f:
            f.write(text)
        sys.stderr.write('已寫入 %s\n' % args.output)
    else:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stdout.write(text)


if __name__ == '__main__':
    main()
